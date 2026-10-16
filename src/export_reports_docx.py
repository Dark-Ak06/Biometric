"""
Скрипт конвертации отчётов по лабораторным работам (Лаб 1, 3, 4, 5, 6)
из формата Markdown (.md) в профессионально оформленные документы Microsoft Word (.docx)
с встроенными графиками, таблицами, кодовыми блоками и математическими формулами.
"""

import os
import re
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls


PROJECT_ROOT = Path(__file__).resolve().parent.parent
REPORTS_DIR = PROJECT_ROOT / "reports"


def set_cell_background(cell, fill_hex: str):
    """Задаёт цвет фона ячейки таблицы Word."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    """Задаёт внутренние отступы ячейки таблицы (в dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_paragraph_shading_and_border(paragraph, fill_hex="F1F5F9", border_hex="CBD5E1"):
    """Оформляет абзац как блок кода или заметки с фоном и рамкой."""
    pPr = paragraph._p.get_or_add_pPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    pPr.append(shd)
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="4" w:color="{border_hex}"/>'
        f'<w:left w:val="single" w:sz="12" w:space="8" w:color="3B82F6"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="4" w:color="{border_hex}"/>'
        f'<w:right w:val="single" w:sz="4" w:space="4" w:color="{border_hex}"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)


def clean_latex_math(text: str) -> str:
    """Преобразует LaTeX-формулы и спецсимволы в читаемые математические Unicode-выражения."""
    # Специфичные блочные и сложные формулы (прямая замена)
    exact_formulas = {
        r"$$S(\mathbf{v}_{\text{live}}, \mathbf{v}_{\text{db}}) = \frac{1 + \frac{\mathbf{v}_{\text{live}} \cdot \mathbf{v}_{\text{db}}}{\|\mathbf{v}_{\text{live}}\| \cdot \|\mathbf{v}_{\text{db}}\|}}{2}$$":
            "S(v_live, v_db) = (1 + (v_live · v_db) / (||v_live|| · ||v_db||)) / 2",
        r"$$T^* = \frac{\min(\text{Score}_{\text{gen}}) + \max(\text{Score}_{\text{imp}})}{2} = \frac{0.326 + 0.258}{2} = \mathbf{0.292}$$":
            "T* = (min(Score_gen) + max(Score_imp)) / 2 = (0.326 + 0.258) / 2 = 0.292",
        r"$CN = \frac{1}{2}\sum_{i=1}^8 |P_i - P_{i+1}|$":
            "CN = 0.5 · Σ(i=1..8) |P_i − P_{i+1}|",
    }
    for old_f, new_f in exact_formulas.items():
        text = text.replace(old_f, new_f)

    def replace_math_span(match):
        m = match.group(1)
        m = re.sub(r'\\text\{([^}]*)\}', r'\1', m)
        m = re.sub(r'\\mathbf\{([^}]*)\}', r'\1', m)
        m = re.sub(r'\\mathcal\{C\}', 'C', m)
        m = re.sub(r'\\mathbb\{R\}\^\{([^}]*)\}', r'R^\1', m)
        m = re.sub(r'\\mathbb\{R\}\^d', 'R^d', m)

        replacements = [
            (r'\\cos', 'cos'),
            (r'\\le', '≤'),
            (r'\\ge', '≥'),
            (r'\\neq', '≠'),
            (r'\\approx', '≈'),
            (r'\\sim', '~'),
            (r'\\pm', '±'),
            (r'\\times', '×'),
            (r'\\cdot', '·'),
            (r'\\implies', '⇒'),
            (r'\\to', '→'),
            (r'\\in', '∈'),
            (r'\\cap', '∩'),
            (r'\\subseteq', '⊆'),
            (r'\\emptyset', '∅'),
            (r'\\land', '∧'),
            (r'\\gg', '≫'),
            (r'\\dots', '…'),
            (r'\\mu_1', 'μ₁'),
            (r'\\mu_0', 'μ₀'),
            (r'\\sigma_1', 'σ₁'),
            (r'\\sigma_0', 'σ₀'),
            (r'\\sigma', 'σ'),
            (r'\\mu', 'μ'),
            (r'\\theta', 'θ'),
            (r'\^\\circ', '°'),
            (r'\^\*', '*'),
            (r'_\{1:N\}', ' (1:N)'),
            (r'_\{1:1\}', ' (1:1)'),
            (r'_\{([^}]*)\}', r'_\1'),
            (r'\\_', '_'),
            (r'\\\%', '%'),
        ]
        for pat, rep in replacements:
            m = re.sub(pat, rep, m)
        return m

    text = re.sub(r'\$\$([^$]+)\$\$', replace_math_span, text)
    text = re.sub(r'\$([^$]+)\$', replace_math_span, text)
    return text


def add_formatted_text(paragraph, text: str, base_font="Times New Roman", base_size=12):
    """
    Парсит inline-разметку Markdown (**жирный**, *курсив*, `код`)
    и добавляет отформатированные Run в абзац Word.
    """
    text = clean_latex_math(text)
    # Токенизация по **bold**, `code`, *italic*
    token_pattern = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)')
    pos = 0
    for match in token_pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            run.font.name = base_font
            run.font.size = Pt(base_size)
        token = match.group(0)
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.font.name = base_font
            run.font.size = Pt(base_size)
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(max(base_size - 1.5, 9))
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        elif token.startswith('*') and token.endswith('*'):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            run.font.name = base_font
            run.font.size = Pt(base_size)
        pos = match.end()

    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        run.font.name = base_font
        run.font.size = Pt(base_size)


def configure_document_styles(doc: Document):
    """Настраивает базовые поля страницы и шрифты документа Word."""
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(0.7)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0x11, 0x11, 0x11)


def render_markdown_to_doc(doc: Document, md_path: Path):
    """Рендерит содержимое одного Markdown-отчёта в объект Document."""
    lines = md_path.read_text(encoding='utf-8').splitlines()
    i = 0
    n = len(lines)

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # Пустая строка или горизонтальный разделитель
        if not stripped or stripped == '---':
            i += 1
            continue

        # Блок кода (``` ... ```)
        if stripped.startswith('```'):
            lang = stripped[3:].strip().lower()
            code_lines = []
            i += 1
            while i < n and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            i += 1  # Пропускаем закрывающие ```

            # Если это mermaid-схема, преобразуем её в аккуратный текстовый блок потока
            if lang == 'mermaid':
                flow_steps = []
                for cl in code_lines:
                    cl_s = cl.strip()
                    if cl_s.startswith('flowchart'):
                        continue
                    cl_s = cl_s.replace('<br/>', ' ')
                    # Упрощаем синтаксис узлов mermaid до читаемой строки
                    cl_s = re.sub(r'[A-Za-z0-9_]+\[\("([^"]+)"\)\]', r'[\1]', cl_s)
                    cl_s = re.sub(r'[A-Za-z0-9_]+\["([^"]+)"\]', r'[\1]', cl_s)
                    cl_s = re.sub(r'[A-Za-z0-9_]+\{"([^"]+)"\}', r'{\1}', cl_s)
                    cl_s = cl_s.replace('-.->', ' ──> ').replace('-->', ' ──> ')
                    cl_s = re.sub(r'\|([^|]+)\|', r' (\1) ', cl_s)
                    flow_steps.append(cl_s)
                code_text = "Логическая схема потока данных:\n" + "\n".join(f"  • {s}" for s in flow_steps)
            else:
                code_text = "\n".join(code_lines)

            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.left_indent = Inches(0.15)
            p.paragraph_format.right_indent = Inches(0.15)
            set_paragraph_shading_and_border(p, fill_hex="F8FAFC", border_hex="CBD5E1")
            run = p.add_run(code_text)
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            continue

        # Таблица Markdown (| ... |)
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines = []
            while i < n and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1

            # Парсим строки таблицы, пропуская строку-разделитель |---|---|
            parsed_rows = []
            for tl in table_lines:
                cells = [c.strip() for c in tl.strip('|').split('|')]
                if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                    continue
                parsed_rows.append(cells)

            if parsed_rows:
                num_cols = max(len(r) for r in parsed_rows)
                table = doc.add_table(rows=len(parsed_rows), cols=num_cols)
                table.style = 'Table Grid'
                table.alignment = WD_TABLE_ALIGNMENT.CENTER

                for r_idx, row_data in enumerate(parsed_rows):
                    row = table.rows[r_idx]
                    for c_idx in range(num_cols):
                        cell = row.cells[c_idx]
                        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                        set_cell_margins(cell, top=80, bottom=80, left=110, right=110)
                        cell_text = row_data[c_idx] if c_idx < len(row_data) else ""
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)

                        if r_idx == 0:
                            set_cell_background(cell, "1E293B")
                            clean_hdr = clean_latex_math(cell_text.replace('**', '').replace('`', ''))
                            run = p.add_run(clean_hdr)
                            run.bold = True
                            run.font.name = "Times New Roman"
                            run.font.size = Pt(10.5)
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        else:
                            if r_idx % 2 == 0:
                                set_cell_background(cell, "F8FAFC")
                            add_formatted_text(p, cell_text, base_font="Times New Roman", base_size=10.5)

                spacer = doc.add_paragraph()
                spacer.paragraph_format.space_before = Pt(2)
                spacer.paragraph_format.space_after = Pt(4)
            continue

        # Изображение (![alt](path))
        img_match = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)$', stripped)
        if img_match:
            alt_text, img_rel_path = img_match.group(1), img_match.group(2)
            img_path = (md_path.parent / img_rel_path).resolve()
            if not img_path.exists():
                img_path = (PROJECT_ROOT / img_rel_path).resolve()

            if img_path.exists():
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(2)
                run_img = p_img.add_run()
                # Оптимальная ширина для страницы A4/Letter
                width_in = 4.6 if "camera_test" in img_path.name else 6.3
                run_img.add_picture(str(img_path), width=Inches(width_in))

                if alt_text:
                    p_cap = doc.add_paragraph()
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cap.paragraph_format.space_before = Pt(2)
                    p_cap.paragraph_format.space_after = Pt(8)
                    r_cap = p_cap.add_run(f"Рисунок — {alt_text}")
                    r_cap.italic = True
                    r_cap.font.name = "Times New Roman"
                    r_cap.font.size = Pt(10.5)
                    r_cap.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            i += 1
            continue

        # Заголовки (#, ##, ###, ####)
        if stripped.startswith('#'):
            hashes = len(stripped) - len(stripped.lstrip('#'))
            title_text = stripped[hashes:].strip()
            level = min(hashes, 4)
            h_para = doc.add_heading(level=level)
            h_para.paragraph_format.space_before = Pt(12 if level <= 2 else 8)
            h_para.paragraph_format.space_after = Pt(4)
            if level == 1:
                h_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

            sizes = {1: 16, 2: 14, 3: 12.5, 4: 12}
            add_formatted_text(h_para, title_text, base_font="Times New Roman", base_size=sizes.get(level, 12))
            for r in h_para.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            i += 1
            continue

        # Цитата (> ...)
        if stripped.startswith('>'):
            quote_text = stripped.lstrip('>').strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.right_indent = Inches(0.15)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            set_paragraph_shading_and_border(p, fill_hex="EFF6FF", border_hex="BFDBFE")
            add_formatted_text(p, quote_text, base_font="Times New Roman", base_size=11.5)
            i += 1
            continue

        # Маркированный или нумерованный список
        bullet_match = re.match(r'^(\s*)([-*]|\d+\.)\s+(.*)$', line)
        if bullet_match:
            indent_spaces = len(bullet_match.group(1))
            marker = bullet_match.group(2)
            item_text = bullet_match.group(3).strip()

            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.left_indent = Inches(0.25 + (0.20 if indent_spaces >= 2 else 0.0))

            prefix = "• " if marker in ('-', '*') else f"{marker} "
            r_pre = p.add_run(prefix)
            r_pre.bold = (marker not in ('-', '*'))
            r_pre.font.name = "Times New Roman"
            r_pre.font.size = Pt(12)

            add_formatted_text(p, item_text, base_font="Times New Roman", base_size=12)
            i += 1
            continue

        # Обычный абзац текста
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        add_formatted_text(p, stripped, base_font="Times New Roman", base_size=12)
        i += 1


def main():
    report_files = [
        "lab01_report.md",
        "lab03_report.md",
        "lab04_report.md",
        "lab05_report.md",
        "lab06_report.md",
    ]

    generated_docx = []

    # 1. Создаём отдельные .docx файлы для каждой лабораторной работы
    for md_name in report_files:
        md_path = REPORTS_DIR / md_name
        if not md_path.exists():
            print(f"[WARN] Файл не найден: {md_path}")
            continue

        docx_path = REPORTS_DIR / md_name.replace(".md", ".docx")
        doc = Document()
        configure_document_styles(doc)
        render_markdown_to_doc(doc, md_path)
        doc.save(str(docx_path))
        size_kb = docx_path.stat().st_size / 1024.0
        print(f"[OK] Сформирован: {docx_path.name} ({size_kb:.1f} КБ)")
        generated_docx.append(docx_path)

    # 2. Создаём единый сводный .docx со всеми 5 лабораторными работами
    combined_path = REPORTS_DIR / "Biometrics_All_Labs_Report.docx"
    combined_doc = Document()
    configure_document_styles(combined_doc)

    for idx, md_name in enumerate(report_files):
        md_path = REPORTS_DIR / md_name
        if not md_path.exists():
            continue
        if idx > 0:
            combined_doc.add_page_break()
        render_markdown_to_doc(combined_doc, md_path)

    combined_doc.save(str(combined_path))
    print(f"[OK] Сформирован сводный отчёт: {combined_path.name} ({combined_path.stat().st_size / 1024.0:.1f} КБ)")


if __name__ == "__main__":
    main()
